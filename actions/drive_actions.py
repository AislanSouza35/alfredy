"""
Leitura do conteúdo dos arquivos que os alunos entregam no Drive.

Antes, o ALF via a entrega e dizia apenas "tem um arquivo chamado
trabalho.pdf, abra para ver". Isso deixava a correção pela metade: ele
sabia quem entregou, mas não o que foi entregue.

Aqui o arquivo é baixado e transformado em texto, de acordo com o tipo:

    Google Docs / Slides   -> exportado como texto
    Google Sheets          -> exportado como CSV
    PDF                    -> texto extraído página a página
    .docx                  -> parágrafos e tabelas
    .txt, .md, código      -> lido direto
    imagem                 -> descrita pelo Gemini Vision

A leitura é somente leitura: nada no Drive do aluno é alterado.
"""

import io
import unicodedata

from core.config import GEMINI_API_KEY


# Quanto texto no máximo volta de um arquivo. Passar disso enche a
# instrução do modelo sem ajudar a corrigir.
MAXIMO_TEXTO = 20_000

# Arquivo maior que isto não é baixado. Trabalho de aluno raramente
# passa disso, e baixar 100 MB travaria a conversa.
MAXIMO_DOWNLOAD = 25 * 1024 * 1024

# Como cada tipo do Google é exportado.
EXPORTACOES_GOOGLE = {
    "application/vnd.google-apps.document": "text/plain",
    "application/vnd.google-apps.presentation": "text/plain",
    "application/vnd.google-apps.spreadsheet": "text/csv",
}

TIPOS_DE_IMAGEM = (
    "image/png",
    "image/jpeg",
    "image/jpg",
    "image/webp",
    "image/gif",
)

TIPOS_DE_TEXTO_PURO = (
    "text/",
    "application/json",
    "application/xml",
    "application/javascript",
)


def _cortar(texto):
    texto = (texto or "").strip()

    if len(texto) > MAXIMO_TEXTO:
        return (
            texto[:MAXIMO_TEXTO]
            + "\n\n[...texto cortado por ser muito longo...]"
        )

    return texto


def obter_servico_drive(credenciais=None):
    """
    Devolve (servico, erro) para ler arquivos do Drive.

    As credenciais vêm da conta dona da turma. Com mais de uma conta
    autorizada isso é obrigatório: o arquivo do aluno de uma turma da
    conta institucional não é visível pela conta pessoal.

    Sem credenciais informadas, usa a primeira conta autorizada.
    """

    if credenciais is None:
        from actions.classroom_contas import servicos

        contas = servicos()
        if not contas:
            return None, (
                "Nenhuma conta do Google está autorizada ainda."
            )
        credenciais = contas[0][2]

    try:
        from googleapiclient.discovery import build

        return build(
            "drive", "v3", credentials=credenciais, cache_discovery=False
        ), None
    except Exception as erro:
        return None, f"Não consegui acessar o Drive: {erro}"


# ============================================================
# CONVERSÃO POR TIPO
# ============================================================

def _texto_de_pdf(dados):
    try:
        from pypdf import PdfReader

        leitor = PdfReader(io.BytesIO(dados))
        paginas = []

        for numero, pagina in enumerate(leitor.pages, start=1):
            texto = (pagina.extract_text() or "").strip()
            if texto:
                paginas.append(f"[página {numero}]\n{texto}")

        if not paginas:
            return (
                "O PDF não tem texto selecionável. Provavelmente é um "
                "documento digitalizado ou uma foto. Diga ao usuário que "
                "ele precisa abrir para ver."
            )

        return "\n\n".join(paginas)

    except Exception as erro:
        return f"Não consegui ler o PDF: {erro}"


def _texto_de_docx(dados):
    try:
        import docx

        documento = docx.Document(io.BytesIO(dados))
        partes = [p.text for p in documento.paragraphs if p.text.strip()]

        for tabela in documento.tables:
            for linha in tabela.rows:
                celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
                if celulas:
                    partes.append(" | ".join(celulas))

        if not partes:
            return "O documento está vazio."

        return "\n".join(partes)

    except Exception as erro:
        return f"Não consegui ler o documento do Word: {erro}"


def _texto_de_xlsx(dados):
    try:
        import openpyxl

        planilha = openpyxl.load_workbook(io.BytesIO(dados), data_only=True)
        partes = []

        for aba in planilha.worksheets:
            partes.append(f"[aba {aba.title}]")
            for linha in aba.iter_rows(max_row=200, values_only=True):
                valores = [
                    str(v) for v in linha if v is not None and str(v).strip()
                ]
                if valores:
                    partes.append(" | ".join(valores))

        return "\n".join(partes) if partes else "A planilha está vazia."

    except Exception as erro:
        return f"Não consegui ler a planilha: {erro}"


def _descrever_imagem(dados, nome):
    """
    Manda a imagem ao Gemini Vision para virar texto.

    Muito aluno entrega foto do caderno ou print de tela; sem isto a
    entrega ficaria ilegível para o ALF.
    """

    if not GEMINI_API_KEY:
        return "Não consigo analisar a imagem: falta a chave da API."

    try:
        from google import genai
        from google.genai import types

        from core.gemini_ssl import criar_http_options_gemini

        cliente = genai.Client(
            api_key=GEMINI_API_KEY,
            http_options=criar_http_options_gemini(types),
        )

        resposta = cliente.models.generate_content(
            model="gemini-3.1-flash-lite",
            contents=[
                (
                    "Esta é a entrega de um aluno, enviada como imagem. "
                    "Transcreva fielmente todo o texto que aparece, na "
                    "ordem em que está. Se houver diagrama, desenho ou "
                    "código, descreva o que mostra. Não avalie e não dê "
                    "nota: apenas relate o conteúdo."
                ),
                types.Part.from_bytes(data=dados, mime_type="image/jpeg"),
            ],
        )

        return (resposta.text or "").strip() or "Não consegui ler a imagem."

    except Exception as erro:
        return f"Não consegui analisar a imagem {nome}: {erro}"


def _converter(dados, tipo, nome):
    tipo = (tipo or "").lower()

    if tipo == "application/pdf" or nome.lower().endswith(".pdf"):
        return _texto_de_pdf(dados)

    if nome.lower().endswith(".docx") or "wordprocessingml" in tipo:
        return _texto_de_docx(dados)

    if nome.lower().endswith((".xlsx", ".xlsm")) or "spreadsheetml" in tipo:
        return _texto_de_xlsx(dados)

    if tipo.startswith(TIPOS_DE_IMAGEM):
        return _descrever_imagem(dados, nome)

    if tipo.startswith(TIPOS_DE_TEXTO_PURO):
        try:
            return dados.decode("utf-8")
        except UnicodeDecodeError:
            return dados.decode("latin-1", errors="replace")

    # Última tentativa: muitos arquivos de código chegam sem tipo certo.
    try:
        return dados.decode("utf-8")
    except UnicodeDecodeError:
        return (
            f"O arquivo {nome} é de um tipo que não consigo ler "
            f"({tipo or 'desconhecido'}). Diga ao usuário que ele precisa "
            "abrir para ver."
        )


# ============================================================
# LEITURA
# ============================================================

def ler_arquivo_do_drive(id_arquivo, nome_sugerido="", credenciais=None):
    """
    Baixa um arquivo do Drive e devolve o conteúdo como texto.
    """

    servico, erro = obter_servico_drive(credenciais)
    if erro:
        return erro

    try:
        from googleapiclient.http import MediaIoBaseDownload
    except ImportError:
        return "As bibliotecas do Google Drive não estão instaladas."

    try:
        info = servico.files().get(
            fileId=id_arquivo, fields="name,mimeType,size"
        ).execute()
    except Exception as erro:
        texto = str(erro)
        if "404" in texto:
            return (
                "Não consegui abrir esse arquivo. Ele pode ter sido "
                "apagado, ou o aluno pode não ter dado acesso a você."
            )
        if "403" in texto:
            return (
                "O Google recusou o acesso a esse arquivo. Confira se "
                "você tem permissão para vê-lo."
            )
        return f"Não consegui abrir o arquivo: {erro}"

    nome = info.get("name", nome_sugerido or "arquivo")
    tipo = info.get("mimeType", "")

    tamanho = info.get("size")
    if tamanho and int(tamanho) > MAXIMO_DOWNLOAD:
        return (
            f"{nome} tem {int(tamanho) // (1024 * 1024)} MB, grande demais "
            "para eu abrir. Diga ao usuário que ele precisa ver direto."
        )

    try:
        if tipo in EXPORTACOES_GOOGLE:
            # Arquivo nativo do Google não se baixa: se exporta.
            requisicao = servico.files().export_media(
                fileId=id_arquivo, mimeType=EXPORTACOES_GOOGLE[tipo]
            )
        else:
            requisicao = servico.files().get_media(fileId=id_arquivo)

        buffer = io.BytesIO()
        download = MediaIoBaseDownload(buffer, requisicao)

        concluido = False
        while not concluido:
            _, concluido = download.next_chunk()

        dados = buffer.getvalue()

    except Exception as erro:
        return f"Não consegui baixar o arquivo {nome}: {erro}"

    if not dados:
        return f"O arquivo {nome} está vazio."

    if tipo in EXPORTACOES_GOOGLE:
        try:
            conteudo = dados.decode("utf-8")
        except UnicodeDecodeError:
            conteudo = dados.decode("latin-1", errors="replace")
    else:
        conteudo = _converter(dados, tipo, nome)

    return f"Conteúdo de {nome}:\n\n{_cortar(conteudo)}"
