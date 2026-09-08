import unicodedata
from pathlib import Path

import openpyxl

from actions.file_actions import area_de_trabalho

LIMITE_ALUNOS_EXIBIDOS = 60


def _normalizar(texto):
    texto = str(texto or "").strip().upper()
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(caractere for caractere in texto if unicodedata.category(caractere) != "Mn")
    return " ".join(texto.split())


def _pastas_busca():
    pastas = [area_de_trabalho()]
    home = Path.home()
    for extra in (home / "Downloads", home / "OneDrive" / "Downloads"):
        if extra.exists() and extra not in pastas:
            pastas.append(extra)
    return pastas


def _localizar_planilha(nome_arquivo):
    nome_arquivo = str(nome_arquivo or "").strip()
    if not nome_arquivo:
        return None, "Informe o nome do arquivo de planilha."

    alvo = _normalizar(nome_arquivo)
    candidatos = []
    for pasta in _pastas_busca():
        for item in pasta.rglob("*.xlsx"):
            if alvo in _normalizar(item.stem) or _normalizar(item.name) == alvo:
                candidatos.append(item)

    if not candidatos:
        return None, f"Não encontrei nenhuma planilha .xlsx chamada '{nome_arquivo}' na Área de Trabalho ou Downloads."
    if len(candidatos) > 1:
        nomes = ", ".join(item.name for item in candidatos[:8])
        return None, f"Encontrei mais de uma planilha parecida: {nomes}. Informe o nome completo."
    return candidatos[0], None


def _selecionar_aba(pasta_trabalho, aba):
    aba = str(aba or "").strip()
    if not aba:
        return pasta_trabalho.active, None

    aba_normalizada = _normalizar(aba)
    for nome_aba in pasta_trabalho.sheetnames:
        if _normalizar(nome_aba) == aba_normalizada or aba_normalizada in _normalizar(nome_aba):
            return pasta_trabalho[nome_aba], None

    abas_disponiveis = ", ".join(pasta_trabalho.sheetnames)
    return None, f"Não encontrei a aba '{aba}'. Abas disponíveis: {abas_disponiveis}."


def ler_planilha_notas(nome_arquivo, aba=""):
    """
    Localiza e lê uma planilha .xlsx de notas na Área de Trabalho ou Downloads.
    Identifica automaticamente a coluna do nome do aluno e devolve, em texto,
    o nome de cada aluno com os valores das demais colunas (notas/avaliações).
    """
    caminho, erro = _localizar_planilha(nome_arquivo)
    if erro:
        return erro

    try:
        pasta_trabalho = openpyxl.load_workbook(caminho, data_only=True, read_only=True)
    except Exception as erro_leitura:
        return f"Não consegui abrir a planilha: {erro_leitura}"

    planilha, erro_aba = _selecionar_aba(pasta_trabalho, aba)
    if erro_aba:
        return erro_aba

    linhas = list(planilha.iter_rows(values_only=True))
    if not linhas:
        return "A planilha está vazia."

    cabecalho = [str(celula or "").strip() for celula in linhas[0]]
    cabecalho_normalizado = [_normalizar(celula) for celula in cabecalho]

    indice_nome = None
    for indice, titulo in enumerate(cabecalho_normalizado):
        if "ALUNO" in titulo or "NOME" in titulo or "ESTUDANTE" in titulo:
            indice_nome = indice
            break

    if indice_nome is None:
        return "Não consegui identificar a coluna com o nome do aluno na planilha."

    colunas_valor = [
        (indice, cabecalho[indice])
        for indice in range(len(cabecalho))
        if indice != indice_nome and cabecalho[indice]
    ]

    if not colunas_valor:
        return "Não encontrei nenhuma coluna de nota na planilha."

    linhas_aluno = []
    for linha in linhas[1:]:
        if indice_nome >= len(linha):
            continue

        nome_aluno = str(linha[indice_nome] or "").strip()
        if not nome_aluno:
            continue

        valores = []
        for indice_coluna, titulo_coluna in colunas_valor:
            valor = linha[indice_coluna] if indice_coluna < len(linha) else None
            valores.append(f"{titulo_coluna}={'' if valor is None else str(valor).strip()}")

        linhas_aluno.append(f"{nome_aluno}: " + "; ".join(valores))

    if not linhas_aluno:
        return "Nenhum aluno com nome preenchido foi encontrado na planilha."

    colunas_texto = ", ".join(titulo for _, titulo in colunas_valor)
    corpo = "\n".join(linhas_aluno[:LIMITE_ALUNOS_EXIBIDOS])
    excedente = len(linhas_aluno) - LIMITE_ALUNOS_EXIBIDOS
    aviso_truncado = "" if excedente <= 0 else f"\n(+{excedente} alunos adicionais não exibidos)"

    return (
        f"Planilha '{caminho.name}' lida com sucesso. "
        f"Colunas encontradas: {colunas_texto}. "
        f"{len(linhas_aluno)} alunos encontrados:\n{corpo}{aviso_truncado}"
    )
